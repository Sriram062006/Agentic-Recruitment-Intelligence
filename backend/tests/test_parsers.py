"""
Tests for Document Parsers (PyMuPDF and python-docx).
Verifies extraction, layout awareness, corruption/empty handling, and unsupported file validation.
"""

import io
import docx
import fitz
import pytest

from app.parsers.docx_parser import DOCXParser
from app.parsers.factory import parse_document
from app.parsers.pdf_parser import PDFParser


def create_sample_pdf_bytes(text: str) -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), text)
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def create_sample_docx_bytes(paragraphs: list[str], table_data: list[list[str]] = None) -> bytes:
    doc = docx.Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table_data:
        t = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
        for r_idx, row in enumerate(table_data):
            for c_idx, val in enumerate(row):
                t.cell(r_idx, c_idx).text = val
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


def test_pdf_parser_success():
    sample_text = "Senior Python Engineer\n5 years experience in FastAPI and Docker."
    pdf_bytes = create_sample_pdf_bytes(sample_text)

    parser = PDFParser()
    parsed = parser.parse(pdf_bytes, "test_resume.pdf")

    assert "Senior Python Engineer" in parsed.raw_text
    assert "FastAPI" in parsed.raw_text
    assert parsed.file_name == "test_resume.pdf"
    assert len(parsed.file_hash) == 64
    assert parsed.page_or_section_count >= 1


def test_docx_parser_success():
    paragraphs = ["Lead Data Scientist", "Proficient in PyTorch, NLP, and Machine Learning."]
    table_data = [["Skill", "Years"], ["Python", "6"], ["FAISS", "3"]]
    docx_bytes = create_sample_docx_bytes(paragraphs, table_data)

    parser = DOCXParser()
    parsed = parser.parse(docx_bytes, "profile.docx")

    assert "Lead Data Scientist" in parsed.raw_text
    assert "PyTorch" in parsed.raw_text
    assert "FAISS" in parsed.raw_text
    assert len(parsed.file_hash) == 64


def test_parser_empty_document_raises_error():
    pdf_parser = PDFParser()
    with pytest.raises(ValueError, match="Empty PDF file"):
        pdf_parser.parse(b"", "empty.pdf")

    docx_parser = DOCXParser()
    with pytest.raises(ValueError, match="Empty DOCX file"):
        docx_parser.parse(b"", "empty.docx")


def test_parser_corrupted_document_raises_error():
    corrupted_bytes = b"%PDF-1.4 garbage corrupt binary bytes 12345"
    pdf_parser = PDFParser()
    with pytest.raises(ValueError, match="Malformed"):
        pdf_parser.parse(corrupted_bytes, "bad.pdf")


def test_parser_unsupported_format_raises_error():
    with pytest.raises(ValueError, match="Unsupported file format"):
        parse_document(b"image binary bytes", "photo.png")
