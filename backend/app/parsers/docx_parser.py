"""
DOCX Document Parser using python-docx.
Extracts structured paragraphs and tables from Word documents.
"""

import hashlib
import io
from docx import Document

from app.core.logging import get_logger
from app.parsers.base import BaseDocumentParser, ParsedDocument

logger = get_logger("docx_parser")


class DOCXParser(BaseDocumentParser):
    def parse(self, content: bytes, file_name: str) -> ParsedDocument:
        if not content or len(content.strip()) == 0:
            raise ValueError(f"Empty DOCX file: '{file_name}' contains 0 bytes")

        file_hash = hashlib.sha256(content).hexdigest()

        try:
            doc = Document(io.BytesIO(content))
        except Exception as e:
            raise ValueError(f"Malformed or unreadable DOCX file '{file_name}': {e}") from e

        extracted_lines = []

        # 1. Paragraphs
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                extracted_lines.append(text)

        # 2. Tables
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    extracted_lines.append(" | ".join(row_cells))

        raw_text = "\n\n".join(extracted_lines).strip()
        section_count = len(doc.sections)

        return ParsedDocument(
            raw_text=raw_text,
            file_name=file_name,
            file_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            file_hash=file_hash,
            page_or_section_count=section_count,
            metadata={"paragraphs": str(len(doc.paragraphs)), "tables": str(len(doc.tables))},
        )
